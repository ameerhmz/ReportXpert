from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from .office_styler import (
    ThemeColors,
    add_executive_cover_page,
    configure_document_headers_footers,
    style_table_element,
    add_callout_box,
    add_official_sign_off_matrix,
    convert_markdown_to_rich_docx
)
from ..core.org_config import org_config


def create_naac_approval_note(
    output_path: Path,
    subject: str,
    reference_no: str,
    department_id: str,
    findings: str,
    audit_verdict: str,
    components_list: list,
    action_items: list
) -> str:
    """
    Generates an executive NAAC/UGC Compliance Approval Note in Microsoft Word (.docx)
    with cover page, headers/footers, styled tables, callout boxes, and sign-off.
    """
    doc = Document()

    # 1. Executive Cover Page
    add_executive_cover_page(
        doc=doc,
        title=f"EXECUTIVE APPROVAL NOTE: {subject.upper()}",
        subtitle=f"Statutory Academic Quality Verification for {department_id}",
        framework="NAAC / UGC",
        department=department_id,
        reference_no=reference_no
    )

    # 2. Running Headers & Footers
    configure_document_headers_footers(doc, framework="NAAC")

    # 3. Section 1: Executive Summary & Intent
    h1 = doc.add_heading("1. Executive Summary & Administrative Intent", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p_exec = doc.add_paragraph(
        f"This approval note conveys the automated statutory compliance audit conducted on "
        f"department '{department_id}' under NAAC Assessment Guidelines. Grounded in the on-premise "
        f"University Knowledge Vault, this document consolidates verified faculty research, student outcomes, "
        f"and statutory procedures. Items flagged below require review prior to final IQAC sign-off."
    )
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(12)

    # 4. Section 2: Component Inventory Table
    h2 = doc.add_heading("2. Extracted Academic Document & Evidence Inventory", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Document / Component Tag"
    hdr_cells[1].text = "Classification Category"
    hdr_cells[2].text = "Verification Status"

    for comp in components_list:
        row_cells = table.add_row().cells
        row_cells[0].text = str(comp.get("tag", "N/A"))
        row_cells[1].text = str(comp.get("type", "Statutory Document"))
        row_cells[2].text = str(comp.get("status", "Verified"))

    style_table_element(table, header_bg=ThemeColors.NAVY_DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 5. Section 3: Statutory Audit Findings
    h3 = doc.add_heading("3. Statutory NAAC / UGC Quality Audit Findings", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    # Convert findings with markdown support
    convert_markdown_to_rich_docx(doc, findings, framework="NAAC")

    # Verdict Box
    callout_type = "success" if "COMPLIANCE VERIFIED" in audit_verdict.upper() else "warning"
    add_callout_box(
        doc,
        f"{audit_verdict}\n\nAutomated analysis executed inside ReportXpert sovereign sandbox. Zero cloud data egress.",
        callout_type=callout_type
    )

    # 6. Section 4: Mandatory Action Items
    h4 = doc.add_heading("4. Mandatory Remediation & Action Plan Before Submission", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    for idx, act in enumerate(action_items, 1):
        p_act = doc.add_paragraph()
        p_act.paragraph_format.space_before = Pt(2)
        p_act.paragraph_format.space_after = Pt(3)
        r_num = p_act.add_run(f"[{idx}] ")
        r_num.bold = True
        r_num.font.name = "Arial"
        r_num.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        r_act = p_act.add_run(act)
        r_act.font.name = "Arial"
        r_act.font.size = Pt(10)

    # 7. Official Sign-Off Block
    add_official_sign_off_matrix(
        doc,
        left=f"Head of Department ({department_id})",
        center="Director (IQAC Secretariat)",
        right="Dean & Vice-Chancellor"
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return str(output_path)


def create_formal_report(
    output_path: Path,
    subject: str,
    reference_no: str,
    line_id: str,
    findings: str,
    audit_verdict: str,
    components_list: list,
    action_items: list
) -> str:
    """Universal formal report generator alias for backward compatibility."""
    return create_naac_approval_note(
        output_path=output_path,
        subject=subject,
        reference_no=reference_no,
        department_id=line_id,
        findings=findings,
        audit_verdict=audit_verdict,
        components_list=components_list,
        action_items=action_items
    )
