from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .core.org_config import org_config

def create_academic_report_sheet(
    output_path: Path,
    department_id: str,
    total_faculty: float,
    total_publications: float,
    total_citations: float,
    h_index_avg: float,
    research_grants_lakhs: float,
    naac_score: float
) -> str:
    """
    Generates an automated Excel calculation sheet for NAAC Department Reports.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NAAC Report"

    # Title Block
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = f"{org_config.org_name.upper()} — ACADEMIC RESEARCH REPORT"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="102C57", end_color="102C57", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Headers
    headers = [
        "Department Code",
        "Total Faculty",
        "Publications",
        "Citations",
        "Avg H-Index",
        "Grants (Lakhs)",
        "Pubs per Faculty",
        "NAAC Score (Proxy)"
    ]

    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 25

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_num)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Populate Data with Live Formulas
    row_num = 4
    
    # Formula: Pubs per Faculty = Publications / Faculty
    f_rate = f"=ROUND(C{row_num}/B{row_num}, 2)"
    f_score = naac_score

    row_vals = [department_id, total_faculty, total_publications, total_citations, h_index_avg, research_grants_lakhs, f_rate, f_score]
    ws.append(row_vals)
    ws.row_dimensions[row_num].height = 20

    for col_num in range(1, 9):
        c = ws.cell(row=row_num, column=col_num)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # Adjust Column Widths
    col_widths = {"A": 20, "B": 15, "C": 15, "D": 15, "E": 14, "F": 16, "G": 20, "H": 22}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return str(output_path)
